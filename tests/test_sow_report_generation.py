"""SOW report families as delivered documents, PM release control, the package.

What is pinned here:

  * a Task 3 weekly / final report is a STRATIFIED LIST of Participants and
    Subparticipants under the four Government categories — entity rows in the
    right table, the Government's words as headings, contract number cited;
  * a record enters a category only on a standing QA approval; pending rows are
    listed separately and never as findings;
  * the weekly report carries suggested changes only, the final report carries
    suggested AND implemented changes (Section C, Task 3);
  * the delivered document passes the automated accessibility checks;
  * PM release moves DRAFT -> PM_REVIEWED -> READY_FOR_DELIVERY and refuses
    anything else, without touching the snapshot or dataset;
  * the package contains the HTML, the CSV, a README with the contract number
    and the SHA-256 of every member.

DEVELOPMENT/TEST DATA. Nothing here is an ONC finding.
"""
from __future__ import annotations

import io
import json
import re
import zipfile

import pytest


class _FakeResult:
    def scalars(self): return self
    def all(self): return []
    def scalar(self): return None
    def scalar_one_or_none(self): return None


class FakeDB:
    """Answers every query empty — the insufficient-data path end to end."""
    async def execute(self, *a, **k): return _FakeResult()
    async def get(self, *a, **k): return None
    def add(self, *a, **k): pass
    async def commit(self): pass
    async def rollback(self): pass


SYNTHETIC_ROWS = {
    "no_discrepancy": [
        {"review_id": "REV-9999-000001", "entity_name": "SYNTHETIC CLINIC ALPHA",
         "entity_level": "Subparticipant", "rce_org_oid": "9.99.555.4.1",
         "qhin": "QHIN 9.99.555.0.1", "category": "no_discrepancy",
         "category_label": "No discrepancies identified", "category_number": 1,
         "rule": "RULE-001 v2", "reportable_at": "2026-09-08T10:00:00",
         "reviewed_at": "2026-09-08T09:00:00"},
    ],
    "minor_administrative": [],
    "inexplicable": [
        {"review_id": "REV-9999-000002", "entity_name": "SYNTHETIC CLINIC BRAVO",
         "entity_level": "Participant", "rce_org_oid": "9.99.555.4.2",
         "qhin": "QHIN 9.99.555.0.1", "category": "inexplicable",
         "category_label": "Inexplicable discrepancies", "category_number": 3,
         "rule": "RULE-004 v2", "reportable_at": "2026-09-09T10:00:00",
         "reviewed_at": "2026-09-09T09:00:00"},
    ],
    "non_compliant": [],
}
SYNTHETIC_PENDING = [
    {"review_id": "REV-9999-000003", "entity_name": "SYNTHETIC CLINIC CHARLIE",
     "entity_level": "Subparticipant", "rce_org_oid": "9.99.555.4.3",
     "qhin": "Unresolved", "category": "inexplicable",
     "category_label": "Inexplicable discrepancies", "category_number": 3,
     "rule": "RULE-004 v2", "reportable_at": None, "reviewed_at": None,
     "pending_reason": "No standing QA approval"},
]


@pytest.fixture
def populated_sow(monkeypatch):
    """Point the stratified list at a fixed synthetic set; leave the rest real."""
    from app.reports.data import sow_report_data as m

    async def _lists(self, review_cycle_id=None, period_start=None, period_end=None):
        return {
            "columns": [{"key": k, "label": v} for k, v in m.SowReportDataService.ENTITY_COLUMNS],
            "entity_lists": {k: list(v) for k, v in SYNTHETIC_ROWS.items()},
            "pending_qa": list(SYNTHETIC_PENDING),
            "counts": {**{k: len(v) for k, v in SYNTHETIC_ROWS.items()},
                       "listed_total": 2, "pending_qa": 1},
            "note": "synthetic",
        }

    monkeypatch.setattr(m.SowReportDataService, "stratified_entities", _lists)
    return m


async def _generate(report_type, **params):
    from app.reports.generator import generate_report

    return await generate_report(FakeDB(), report_type=report_type, persist=False,
                                 query_parameters=params or None)


def _body(html: str) -> str:
    return re.sub(r"<style>.*?</style>", "", html, flags=re.DOTALL)


class TestTaskThreeWeekly:

    @pytest.mark.asyncio
    async def test_weekly_is_a_stratified_entity_list_with_government_wording(self, populated_sow):
        result = await _generate("retrospective_weekly",
                                 period_start="2026-09-07", period_end="2026-09-13",
                                 suggested_changes="Clarify D4 address materiality")
        html = _body(result["html"])
        for label in ("No discrepancies identified", "Minor or administrative discrepancies",
                      "Inexplicable discrepancies", "Non-compliant discrepancies"):
            assert label in html
        # The entities are on the page, in the right category section.
        alpha = html.index("SYNTHETIC CLINIC ALPHA")
        bravo = html.index("SYNTHETIC CLINIC BRAVO")
        assert html.index("6.1 No discrepancies identified") < alpha < html.index("6.2 Minor")
        assert html.index("6.3 Inexplicable discrepancies") < bravo < html.index("6.4 Non-compliant")
        # The pending record is listed as pending, not under a category.
        charlie = html.index("SYNTHETIC CLINIC CHARLIE")
        assert charlie > html.index("Reviews pending independent QA")
        assert "No standing QA approval" in html

    @pytest.mark.asyncio
    async def test_weekly_cites_the_contract_and_the_deliverable(self, populated_sow):
        result = await _generate("retrospective_weekly")
        html = result["html"]
        assert "7571MN26F80064" in html
        assert "Deliverable D3.1" in html
        assert "Task 3 Weekly Progress Report" in html
        assert result["dataset"]["deliverable"] == "D3.1"

    @pytest.mark.asyncio
    async def test_weekly_carries_suggested_but_not_implemented_changes(self, populated_sow):
        result = await _generate("retrospective_weekly",
                                 suggested_changes="Line one\nLine two",
                                 implemented_changes="Should not appear")
        html = _body(result["html"])
        assert "Line one" in html and "Line two" in html
        assert "Implemented changes" not in html
        assert "Should not appear" not in html

    @pytest.mark.asyncio
    async def test_weekly_passes_automated_accessibility_checks(self, populated_sow):
        result = await _generate("retrospective_weekly")
        assert result["accessibility"]["automated_checks_passed"], \
            result["accessibility"]["errors"]

    @pytest.mark.asyncio
    async def test_internal_shorthand_is_not_used_as_a_category_heading(self, populated_sow):
        result = await _generate("retrospective_weekly")
        headings = re.findall(r"<h[1-4][^>]*>(.*?)</h[1-4]>", _body(result["html"]), re.DOTALL)
        for h in headings:
            assert not re.search(r"\bB[1-4]\b", h), h

    @pytest.mark.asyncio
    async def test_csv_lists_each_entity_with_its_category_number(self, populated_sow):
        result = await _generate("retrospective_weekly")
        csv_text = result["csv"]
        assert "# Contract number: 7571MN26F80064" in csv_text
        assert "SYNTHETIC CLINIC ALPHA" in csv_text
        line = next(l for l in csv_text.splitlines() if "SYNTHETIC CLINIC BRAVO" in l)
        assert line.startswith("3,")
        assert "SYNTHETIC CLINIC CHARLIE" in csv_text.split("## Pending independent QA")[1]


class TestTaskThreeFinalAndOtherFamilies:

    @pytest.mark.asyncio
    async def test_final_report_carries_implemented_changes(self, populated_sow):
        result = await _generate("retrospective_final",
                                 suggested_changes="S1", implemented_changes="I1")
        html = _body(result["html"])
        assert "Implemented changes" in html and "I1" in html and "S1" in html
        assert result["dataset"]["deliverable"] == "D3.2"

    @pytest.mark.asyncio
    @pytest.mark.parametrize("report_type,deliverable", [
        ("ongoing_biweekly", "D4.1"), ("ongoing_quarterly", "D4.2"),
        ("priority_status", "D5.1"), ("priority_quarterly", "D5.2"),
    ])
    async def test_every_family_renders_and_passes_accessibility(self, populated_sow,
                                                                 report_type, deliverable):
        result = await _generate(report_type)
        assert result["dataset"]["deliverable"] == deliverable
        assert f"Deliverable {deliverable}" in result["html"]
        assert result["accessibility"]["automated_checks_passed"], \
            result["accessibility"]["errors"]

    @pytest.mark.asyncio
    async def test_empty_database_renders_honestly(self):
        """No stub: every list is empty and the document says so."""
        result = await _generate("retrospective_weekly")
        html = _body(result["html"])
        assert "No Participant or Subparticipant is listed under this category" in html
        assert "No review record is pending QA" in html
        assert result["accessibility"]["automated_checks_passed"]

    def test_report_type_keys_fit_the_column(self):
        from app.reports.data.sow_report_data import SOW_REPORT_TYPES

        for key in SOW_REPORT_TYPES:
            assert len(key) <= 20, key  # review_reports.report_type is String(20)

    def test_the_stale_not_implemented_note_is_gone(self):
        import inspect

        from app.reports.data import sow_report_data

        assert "Per-QHIN sample draw is not implemented" not in inspect.getsource(sow_report_data)


class TestReleaseControl:

    def test_default_is_draft(self):
        from app.reports.data.release import STATUS_DRAFT, current_release

        assert current_release(None)["status"] == STATUS_DRAFT
        assert current_release({"snapshot": {}})["status"] == STATUS_DRAFT

    def test_happy_path_and_refusals(self):
        from app.reports.data.release import (ReleaseTransitionError,
                                              apply_transition)

        data = {"snapshot": {"report_id": "X"}, "dataset": {"a": 1}}
        with pytest.raises(ReleaseTransitionError):
            apply_transition(data, action="READY_FOR_DELIVERY", actor="pm")
        data, e1 = apply_transition(data, action="PM_REVIEWED", actor="pm@x", note="ok")
        assert data["release"]["status"] == "PM_REVIEWED"
        data, e2 = apply_transition(data, action="READY_FOR_DELIVERY", actor="pm@x")
        assert data["release"]["status"] == "READY_FOR_DELIVERY"
        assert [h["status"] for h in data["release"]["history"]] == ["PM_REVIEWED", "READY_FOR_DELIVERY"]
        with pytest.raises(ReleaseTransitionError):
            apply_transition(data, action="PM_REVIEWED", actor="pm@x")
        data, _ = apply_transition(data, action="RETURNED_TO_DRAFT", actor="pm@x")
        data, _ = apply_transition(data, action="PM_REVIEWED", actor="pm@x")
        # The frozen record is untouched throughout.
        assert data["snapshot"] == {"report_id": "X"} and data["dataset"] == {"a": 1}
        with pytest.raises(ReleaseTransitionError):
            apply_transition(data, action="SENT", actor="pm@x")

    def test_package_contents_and_hashes(self):
        import hashlib

        from app.reports.data.release import CONTRACT_NUMBER, build_package, current_release

        html = "<html><body>report</body></html>"
        csv_text = "# header\r\n1,REV,X\r\n"
        pkg = build_package(
            report_id="DA-ARC-2026-999", html=html, csv_text=csv_text, pdf_bytes=None,
            snapshot={"data_classification": "DEVELOPMENT_TEST",
                      "generation_timestamp": "2026-09-10T00:00:00+00:00",
                      "generated_by": "pm@x", "data_payload_hash": "abc"},
            release=current_release({"release": {"status": "READY_FOR_DELIVERY", "history": []}}),
            deliverable={"deliverable": "D3.1", "task": "Task 3", "title": "Task 3 Weekly Progress Report"},
            pdf_unavailable_reason="engine absent")
        archive = zipfile.ZipFile(io.BytesIO(pkg["bytes"]))
        names = set(archive.namelist())
        assert names == {"DA-ARC-2026-999.html", "DA-ARC-2026-999.csv", "README.txt", "manifest.json"}
        readme = archive.read("README.txt").decode("utf-8")
        assert CONTRACT_NUMBER in readme
        assert "NOT FOR GOVERNMENT DELIVERY" in readme
        assert "Ready for delivery" in readme
        assert "has not been transmitted" in readme
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["files"]["DA-ARC-2026-999.html"] == hashlib.sha256(html.encode()).hexdigest()
        assert archive.read("DA-ARC-2026-999.csv").startswith(b"\xef\xbb\xbf")

    def test_route_surface_exists(self):
        from app.reports.routes import router

        paths = {(r.path, tuple(sorted(r.methods))) for r in router.routes}
        assert ("/api/reports/{report_id}/release", ("GET",)) in paths
        assert ("/api/reports/{report_id}/release", ("POST",)) in paths
        assert ("/api/reports/{report_id}/package", ("GET",)) in paths
        assert ("/api/reports/{report_id}/docx", ("GET",)) in paths

    def test_package_uses_the_traceable_stem_and_carries_the_docx(self):
        from app.reports.data.release import build_package, current_release

        stem = "7571MN26F80064_Task3_D3.1_Weekly_2026-09-05_2026-09-11_DA-ARC-2026-016"
        pkg = build_package(
            report_id="DA-ARC-2026-016", html="<html></html>", csv_text="a,b\r\n",
            pdf_bytes=b"%PDF-1.7 fake", docx_bytes=b"PK fake docx", stem=stem,
            snapshot={"data_classification": "DEVELOPMENT_TEST",
                      "generation_timestamp": "2026-09-11T00:00:00+00:00",
                      "generated_by": "pm@x", "data_payload_hash": "abc"},
            release=current_release({}),
            deliverable={"deliverable": "D3.1", "task": "Task 3", "title": "Task 3 Weekly Progress Report"})
        assert pkg["filename"] == f"{stem}.zip"
        names = zipfile.ZipFile(io.BytesIO(pkg["bytes"])).namelist()
        # Editable copy first, then the customer rendering, then archive + data.
        assert names[:4] == [f"{stem}.docx", f"{stem}.pdf", f"{stem}.html", f"{stem}.csv"]
        assert {"README.txt", "manifest.json"} <= set(names)


class TestDeliverableIdentity:
    """Cover, document control, and the Government-mark rule."""

    def test_government_mark_is_off_by_default(self, monkeypatch):
        from app.reports import branding as b

        for name in ("REPORT_GOVERNMENT_BRANDING_AUTHORIZED", "REPORT_GOVERNMENT_LOGO_PATH",
                     "REPORT_AGT_LOGO_PATH"):
            monkeypatch.delenv(name, raising=False)
        current = b.current_branding()
        assert current.government_branding_authorized is False
        assert current.government_logo is None
        assert current.prepared_by == "Alliance Global Tech Inc."
        assert current.prepared_for[0] == "U.S. Department of Health and Human Services"
        assert "Office of the National Coordinator" in current.prepared_for[-1]

    def test_flag_alone_does_not_authorize_a_mark(self, monkeypatch):
        """The flag without an approved image file still renders text only —
        there is nothing to place, and the policy needs written approval."""
        from app.reports import branding as b

        monkeypatch.setenv("REPORT_GOVERNMENT_BRANDING_AUTHORIZED", "true")
        monkeypatch.delenv("REPORT_GOVERNMENT_LOGO_PATH", raising=False)
        assert b.current_branding().government_branding_authorized is False

    def test_filename_stem_is_traceable_and_filesystem_safe(self):
        from app.reports.branding import deliverable_filename_stem

        stem = deliverable_filename_stem(
            contract_number="7571MN26F80064", task="Task 3", deliverable="D3.1",
            kind="Weekly", period_start="2026-09-05", period_end="2026-09-11",
            report_id="DA-ARC-2026-016")
        assert stem == "7571MN26F80064_Task3_D3.1_Weekly_2026-09-05_2026-09-11_DA-ARC-2026-016"
        odd = deliverable_filename_stem(contract_number="X/Y", task=None, deliverable=None,
                                        kind="Ad hoc?", period_start=None, period_end=None,
                                        report_id="R 1")
        assert re.fullmatch(r"[A-Za-z0-9._-]+", odd)

    @pytest.mark.asyncio
    async def test_html_has_cover_document_control_and_text_recipient(self, populated_sow, monkeypatch):
        monkeypatch.delenv("REPORT_GOVERNMENT_BRANDING_AUTHORIZED", raising=False)
        result = await _generate("retrospective_weekly",
                                 period_start="2026-09-07", period_end="2026-09-13")
        html = result["html"]
        body = _body(html)
        assert 'class="cover"' in html
        assert "Document Control" in body
        assert "Prepared for" in body
        assert "U.S. Department of Health and Human Services" in body
        assert "Office of the National Coordinator for Health Information Technology" in body
        assert "Alliance Global Tech Inc." in body
        assert 'class="cover-gov-mark"' not in html  # no Government mark placed (CSS rule may exist)
        assert 'name="author" content="Alliance Global Tech Inc."' in html
        assert 'name="dcterms.created"' in html
        assert "Contract 7571MN26F80064" in body
        # Body sections in contract order.
        order = ["1. Executive Summary", "2. Reporting Period and Scope", "3. Review Activity",
                 "4. Population and Sampling Context", "5. QHIN Stratification Summary",
                 "6. Review Results", "7. Significant Observations",
                 "8. Suggested Methodology / Control Framework Changes",
                 "9. Issues, Dependencies and COR Decisions", "10. Supporting Traceability"]
        positions = [body.index(h) for h in order]
        assert positions == sorted(positions)
        # Identity travels in the frozen dataset as text only.
        assert result["dataset"]["branding"]["government_branding_authorized"] is False
        assert "agt_logo" not in result["dataset"]["branding"]


class TestDocxDeliverable:

    @pytest.mark.asyncio
    async def test_docx_has_styles_repeat_headers_fields_and_properties(self, populated_sow, monkeypatch):
        from docx import Document

        from app.reports.branding import current_branding
        from app.reports.data.release import current_release
        from app.reports.engine.docx_engine import render_sow_docx

        monkeypatch.delenv("REPORT_GOVERNMENT_BRANDING_AUTHORIZED", raising=False)
        result = await _generate("retrospective_weekly",
                                 period_start="2026-09-07", period_end="2026-09-13",
                                 suggested_changes="Clarify D4 address materiality")
        dataset = result["dataset"]
        snapshot = {"report_id": "DA-ARC-2026-016", "data_classification": "DEVELOPMENT_TEST",
                    "generation_timestamp": "2026-09-11T00:00:00+00:00", "generated_by": "pm@x",
                    "template_version": "2.0", "source_provenance": {}, "b1_b4_rule_version": "v2"}
        blob = render_sow_docx(dataset, snapshot, current_release({}), current_branding().to_dict())
        doc = Document(io.BytesIO(blob))

        styles = [p.style.name for p in doc.paragraphs]
        assert "Title" in styles and "Heading 1" in styles and "Heading 2" in styles
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "DEVELOPMENT / TEST DATA" in text
        assert "U.S. Department of Health and Human Services" in text
        assert "Authorized Government mark" not in text
        assert "6.1 No discrepancies identified" in text
        assert "6.3 Inexplicable discrepancies" in text
        assert "Clarify D4 address materiality" in text
        assert "Implemented changes" not in text        # weekly: suggested only

        # Every table repeats its header row across pages and names its columns.
        assert doc.tables, "no tables"
        for table in doc.tables:
            first = table.rows[0]._tr
            assert first.xpath("./w:trPr/w:tblHeader"), "header row does not repeat"
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        assert "SYNTHETIC CLINIC ALPHA" in cells and "SYNTHETIC CLINIC BRAVO" in cells
        assert "SYNTHETIC CLINIC CHARLIE" in cells   # pending, in its own table

        xml = doc.element.xml
        assert 'TOC \\o "1-3"' in xml or "TOC \\o &quot;1-3&quot;" in xml
        footer_xml = doc.sections[0].footer._element.xml
        assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
        assert "7571MN26F80064" in doc.sections[0].header.paragraphs[0].text

        core = doc.core_properties
        assert core.title.startswith("Task 3 Weekly Progress Report DA-ARC-2026-016")
        assert core.author == "Alliance Global Tech Inc."
        assert core.language == "en-US"
        assert "7571MN26F80064" in core.keywords


class TestStoredCsv:
    """The register's CSV of a SOW report is the entity list, not figure data."""

    class _Row:
        def __init__(self, report_type, dataset):
            self.report_id = "DA-ARC-2026-777"
            self.report_type = report_type
            self.report_data = {"dataset": dataset,
                                "snapshot": {"generation_timestamp": "2026-09-11T00:00:00+00:00"}}

    def test_sow_report_csv_is_the_stratified_list(self):
        from app.reports.routes import csv_for_stored_report

        dataset = {
            "deliverable": "D3.1", "deliverable_title": "Task 3 Weekly Progress Report",
            "contract_number": "7571MN26F80064", "service_version": "1.0.0",
            "categories": ["no_discrepancy", "minor_administrative", "inexplicable", "non_compliant"],
            "category_numbers": {"inexplicable": 3},
            "entity_columns": [{"key": "review_id", "label": "Case"}, {"key": "entity_name", "label": "Entity"}],
            "entity_lists": {"inexplicable": [{"review_id": "REV-1", "entity_name": "SYNTH"}]},
            "pending_qa": [], "methodology_changes": {"suggested": ["s"], "includes_implemented": False},
        }
        csv_text = csv_for_stored_report(self._Row("retrospective_weekly", dataset))
        assert "# Contract number: 7571MN26F80064" in csv_text
        assert any(line.startswith("3,REV-1,SYNTH") for line in csv_text.splitlines())

    def test_technical_report_csv_still_uses_the_figure_path(self):
        from app.reports.routes import csv_for_stored_report

        csv_text = csv_for_stored_report(self._Row("verification", {"scope": {}, "service_version": "1.0.0"}))
        assert "Scope at a Glance" in csv_text
